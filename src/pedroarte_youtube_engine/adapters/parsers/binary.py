"""Parsers de formatos binários — dependências opcionais.

Estes formatos exigem o extra `parsers`:

    pip install -e ".[parsers]"

Sem ele, o motor **não quebra**: o parser informa a limitação com uma mensagem
acionável, e o `INPUT_DISCOVERY_AGENT` registra o documento como não lido em vez
de abortar a execução.
"""

from __future__ import annotations

import re
from pathlib import Path

from pedroarte_youtube_engine.ports.parsing import ParsedDocument
from pedroarte_youtube_engine.shared.errors import IngestionError
from pedroarte_youtube_engine.shared.text import normalize_whitespace

_HEADING_HINT = re.compile(
    r"^\s*(CAP[ÍI]TULO|CAPITULO|PARTE|ATO)\s+([\dIVXLCDM]+|[A-Za-zÀ-ÿ]+)\s*[-–—:.]?\s*(.*)$",
    re.IGNORECASE,
)


def _missing(package: str, extra: str, path: Path) -> IngestionError:
    return IngestionError(
        f"O pacote '{package}' não está instalado; não é possível ler {path.name}. "
        f'Instale o extra opcional com: pip install -e ".[{extra}]"',
        path=str(path),
        missing_package=package,
    )


class DocxParser:
    """Lê `.docx` usando `python-docx`, tratando estilos de título como capítulos."""

    @property
    def supported_suffixes(self) -> frozenset[str]:
        return frozenset({".docx"})

    def can_parse(self, path: Path) -> bool:
        return path.suffix.lower() in self.supported_suffixes

    def parse(self, path: Path) -> ParsedDocument:
        try:
            import docx  # type: ignore[import-not-found]
        except ImportError as exc:  # pragma: no cover - depende do ambiente
            raise _missing("python-docx", "parsers", path) from exc

        document = docx.Document(str(path))
        sections: list[tuple[str, list[str]]] = []
        body: list[str] = []
        current_title = ""
        current_body: list[str] = []

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            body.append(text)
            style = (paragraph.style.name or "").lower() if paragraph.style else ""
            is_heading = style.startswith("heading") or style.startswith("título")
            if is_heading or _HEADING_HINT.match(text):
                if current_title and current_body:
                    sections.append((current_title, current_body))
                current_title = text
                current_body = []
            else:
                current_body.append(text)

        if current_title and current_body:
            sections.append((current_title, current_body))

        core = document.core_properties
        return ParsedDocument(
            text=normalize_whitespace("\n\n".join(body)),
            title=(core.title or "").strip(),
            author=(core.author or "").strip(),
            sections=tuple(
                (title, normalize_whitespace("\n\n".join(lines))) for title, lines in sections
            ),
        )


class PdfParser:
    """Lê `.pdf` usando `pypdf`.

    A extração de PDF é intrinsecamente imperfeita: o formato descreve
    posicionamento, não estrutura. O parser registra um aviso explícito para que
    o operador saiba que a divisão em capítulos pode exigir revisão.
    """

    @property
    def supported_suffixes(self) -> frozenset[str]:
        return frozenset({".pdf"})

    def can_parse(self, path: Path) -> bool:
        return path.suffix.lower() in self.supported_suffixes

    def parse(self, path: Path) -> ParsedDocument:
        try:
            from pypdf import PdfReader  # type: ignore[import-not-found]
        except ImportError as exc:  # pragma: no cover - depende do ambiente
            raise _missing("pypdf", "parsers", path) from exc

        reader = PdfReader(str(path))
        pages = [page.extract_text() or "" for page in reader.pages]
        text = normalize_whitespace("\n\n".join(pages))

        metadata = reader.metadata or {}
        title = str(metadata.get("/Title", "") or "").strip()
        author = str(metadata.get("/Author", "") or "").strip()

        sections = _sections_from_headings(text)
        warnings = (
            "A extração de PDF preserva o texto, mas não garante a estrutura de "
            "capítulos. Revise `canon/canon_bible.md` antes de gerar os formatos longos.",
        )
        return ParsedDocument(
            text=text,
            title=title,
            author=author,
            sections=sections,
            warnings=warnings,
        )


class EpubParser:
    """Lê `.epub` usando `EbookLib` + `beautifulsoup4`.

    ATENÇÃO: `EbookLib` é distribuído sob AGPL-3.0. Ele só é carregado quando o
    extra `parsers` está instalado e um EPUB é encontrado. Veja `NOTICE`.
    """

    @property
    def supported_suffixes(self) -> frozenset[str]:
        return frozenset({".epub"})

    def can_parse(self, path: Path) -> bool:
        return path.suffix.lower() in self.supported_suffixes

    def parse(self, path: Path) -> ParsedDocument:
        try:
            import ebooklib  # type: ignore[import-not-found]
            from bs4 import BeautifulSoup  # type: ignore[import-not-found]
            from ebooklib import epub  # type: ignore[import-not-found]
        except ImportError as exc:  # pragma: no cover - depende do ambiente
            raise _missing("EbookLib/beautifulsoup4", "parsers", path) from exc

        book = epub.read_epub(str(path))
        sections: list[tuple[str, str]] = []
        chunks: list[str] = []

        for item in book.get_items_of_type(ebooklib.ITEM_DOCUMENT):
            soup = BeautifulSoup(item.get_content(), "html.parser")
            heading_tag = soup.find(["h1", "h2", "h3"])
            heading = heading_tag.get_text(strip=True) if heading_tag else ""
            body = normalize_whitespace(soup.get_text("\n"))
            if not body:
                continue
            chunks.append(body)
            sections.append((heading or f"Seção {len(sections) + 1}", body))

        titles = book.get_metadata("DC", "title")
        authors = book.get_metadata("DC", "creator")
        return ParsedDocument(
            text=normalize_whitespace("\n\n".join(chunks)),
            title=str(titles[0][0]) if titles else "",
            author=str(authors[0][0]) if authors else "",
            sections=tuple(sections),
        )


def _sections_from_headings(text: str) -> tuple[tuple[str, str], ...]:
    lines = text.split("\n")
    sections: list[tuple[str, list[str]]] = []
    current_title = ""
    current_body: list[str] = []
    for line in lines:
        if _HEADING_HINT.match(line):
            if current_title and current_body:
                sections.append((current_title, current_body))
            current_title = line.strip()
            current_body = []
        else:
            current_body.append(line)
    if current_title and current_body:
        sections.append((current_title, current_body))
    return tuple(
        (title, normalize_whitespace("\n".join(body))) for title, body in sections
    )
